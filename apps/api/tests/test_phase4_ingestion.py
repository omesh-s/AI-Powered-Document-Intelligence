from __future__ import annotations

import asyncio
from io import BytesIO
from typing import Any
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app.api.routes import health as health_mod
from app.core.database import AsyncSessionLocal
from app.core.storage import ObjectStorageClient, PresignedUpload
from app.models.document import Chunk, Page, StructuredBlock
from app.services import ingestion_service
from tests.conftest import FAKE_OBJECT_BYTES, FAKE_OBJECT_RAISE_ON, auth_headers


def _assert_error_envelope(body: dict[str, Any]) -> None:
    assert "error" in body
    err = body["error"]
    assert "code" in err
    assert "message" in err
    assert "requestId" in err
    assert "details" in err


class FakeObjectStorage(ObjectStorageClient):
    def __init__(self, store: dict[str, bytes], raise_on: set[str]) -> None:
        self._store = store
        self._raise_on = raise_on

    def presign_put_object(
        self,
        *,
        bucket: str,
        key: str,
        content_type: str,
        expires_in: int,
    ) -> PresignedUpload:
        raise RuntimeError("presign_put_object not implemented in FakeObjectStorage")

    def get_object_bytes(self, *, bucket: str, key: str) -> bytes:
        _ = bucket
        if key in self._raise_on:
            raise RuntimeError("FakeObjectStorage forced failure")
        return self._store.get(key, b"")

    def delete_object(self, *, bucket: str, key: str) -> None:
        _ = (bucket, key)


class CountingOCR(ingestion_service.OCRProvider):
    def __init__(self, text: str) -> None:
        self.called = 0
        self._text = text

    async def ocr_empty_page(self) -> tuple[str, float | None]:
        self.called += 1
        return self._text, 0.99


async def _process_job(job_id: UUID, *, storage: ObjectStorageClient, ocr_provider=None) -> None:
    async with AsyncSessionLocal() as session:
        await ingestion_service.process_ingestion_job(
            session,
            job_id=job_id,
            storage_client=storage,
            ocr_provider=ocr_provider,
        )


def process_job_sync(job_id: UUID, *, ocr_provider=None) -> None:
    storage = FakeObjectStorage(FAKE_OBJECT_BYTES, FAKE_OBJECT_RAISE_ON)
    asyncio.run(_process_job(job_id, storage=storage, ocr_provider=ocr_provider))


def _register(client: TestClient, email: str) -> str:
    reg = client.post("/api/auth/register", json={"email": email, "password": "securepass1"})
    assert reg.status_code == 200
    return reg.json()["tokens"]["access_token"]


def _create_workspace(client: TestClient, token: str, name: str) -> dict[str, Any]:
    ws = client.post("/api/workspaces/", headers=auth_headers(token), json={"name": name})
    assert ws.status_code == 200
    return ws.json()


def _create_document_via_upload_url(
    client: TestClient,
    *,
    token: str,
    workspace_id: UUID,
    filename: str,
    content_type: str,
    file_size: int,
    content_bytes: bytes,
) -> tuple[dict[str, Any], str, UUID]:
    up = client.post(
        "/api/documents/upload-url",
        headers=auth_headers(token),
        json={
            "workspace_id": workspace_id,
            "filename": filename,
            "content_type": content_type,
            "file_size": file_size,
        },
    )
    assert up.status_code == 200
    upload = up.json()

    storage_key = upload["storage_key"]
    FAKE_OBJECT_BYTES[storage_key] = content_bytes

    doc = client.post(
        "/api/documents/",
        headers=auth_headers(token),
        json={
            "workspace_id": workspace_id,
            "filename": filename,
            "content_type": content_type,
            "file_size": file_size,
            "storage_key": storage_key,
            "checksum_sha256": "a" * 64,
        },
    )
    assert doc.status_code == 200
    doc_body = doc.json()
    job_id = UUID(doc_body["ingestion_job"]["id"])
    return doc_body, storage_key, job_id


def test_get_ingestion_jobs_list_authorization_and_pagination(
    client: TestClient,
    unique_email: str,
) -> None:
    # Create two separate users.
    token1 = _register(client, f"u1_{unique_email}")
    token2 = _register(client, f"u2_{unique_email}")

    ws1 = _create_workspace(client, token1, "WS1")
    ws2 = _create_workspace(client, token2, "WS2")

    # Create 3 docs in ws1 (3 jobs).
    for i in range(3):
        _create_document_via_upload_url(
            client,
            token=token1,
            workspace_id=ws1["id"],
            filename=f"d{i}.txt",
            content_type="text/plain",
            file_size=10,
            content_bytes=b"hello",
        )

    # Create 1 doc in ws2 (1 job).
    _create_document_via_upload_url(
        client,
        token=token2,
        workspace_id=ws2["id"],
        filename="other.txt",
        content_type="text/plain",
        file_size=10,
        content_bytes=b"world",
    )

    # Member sees their jobs with pagination.
    r = client.get(
        "/api/ingestion/jobs",
        headers=auth_headers(token1),
        params={"workspace_id": ws1["id"], "page": 1, "size": 2},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["pagination"]["total"] == 3
    assert len(body["items"]) == 2

    # Non-member should see none.
    r2 = client.get(
        "/api/ingestion/jobs",
        headers=auth_headers(token2),
        params={"workspace_id": ws1["id"], "page": 1, "size": 10},
    )
    assert r2.status_code == 200
    assert r2.json()["pagination"]["total"] == 0


def test_get_ingestion_job_detail_forbidden_non_member(
    client: TestClient,
    unique_email: str,
) -> None:
    token1 = _register(client, f"owner_{unique_email}")
    ws1 = _create_workspace(client, token1, "Owned")
    doc_body, _storage_key, job_id = _create_document_via_upload_url(
        client,
        token=token1,
        workspace_id=ws1["id"],
        filename="a.txt",
        content_type="text/plain",
        file_size=10,
        content_bytes=b"hello",
    )
    assert UUID(doc_body["ingestion_job"]["id"]) == job_id

    token2 = _register(client, f"nonmember_{unique_email}")
    resp = client.get(f"/api/ingestion/jobs/{job_id}", headers=auth_headers(token2))
    assert resp.status_code == 403
    _assert_error_envelope(resp.json())


def test_txt_ingestion_success_and_chunk_metadata(
    client: TestClient,
    unique_email: str,
) -> None:
    token = _register(client, f"txt_{unique_email}")
    ws = _create_workspace(client, token, "TXTWS")

    doc_body, _storage_key, job_id = _create_document_via_upload_url(
        client,
        token=token,
        workspace_id=ws["id"],
        filename="hello.txt",
        content_type="text/plain",
        file_size=10,
        content_bytes=b"Hello world\n\nSecond para",
    )
    version_id = UUID(doc_body["latest_version_id"])

    process_job_sync(job_id)

    job_resp = client.get(f"/api/ingestion/jobs/{job_id}", headers=auth_headers(token))
    assert job_resp.status_code == 200
    assert job_resp.json()["job"]["status"] == "succeeded"
    assert job_resp.json()["job"]["stage"] == "indexed"
    assert job_resp.json()["job"]["progress"] == 1.0

    async def _verify() -> None:
        async with AsyncSessionLocal() as session:
            pages_count = (
                await session.execute(
                    select(func.count())
                    .select_from(Page)
                    .where(Page.document_version_id == version_id)
                )
            ).scalar_one()

            chunks = (
                (
                    await session.execute(
                        select(Chunk)
                        .where(Chunk.document_version_id == version_id)
                        .order_by(Chunk.chunk_index)
                    )
                )
                .scalars()
                .all()
            )
            blocks_count = (
                await session.execute(
                    select(func.count())
                    .select_from(StructuredBlock)
                    .join(Page, Page.id == StructuredBlock.page_id)
                    .where(Page.document_version_id == version_id)
                )
            ).scalar_one()

            assert pages_count == 1
            assert blocks_count >= 1
            assert len(chunks) >= 1
            meta = chunks[0].metadata_json
            assert meta["workspace_id"] == str(ws["id"])
            assert meta["document_id"] == str(doc_body["id"])
            assert meta["document_version_id"] == str(version_id)
            assert meta["page_number"] == 1

    asyncio.run(_verify())


def test_txt_ocr_fallback_called_for_empty_text(
    client: TestClient,
    unique_email: str,
) -> None:
    token = _register(client, f"ocr_{unique_email}")
    ws = _create_workspace(client, token, "OCRWS")

    doc_body, storage_key, job_id = _create_document_via_upload_url(
        client,
        token=token,
        workspace_id=ws["id"],
        filename="empty.txt",
        content_type="text/plain",
        file_size=10,
        content_bytes=b"",
    )
    counter = CountingOCR(text="OCR_FALLBACK_TEXT")
    process_job_sync(job_id, ocr_provider=counter)
    assert counter.called == 1

    async def _verify_chunk_text_contains() -> None:
        async with AsyncSessionLocal() as session:
            version_id = UUID(doc_body["latest_version_id"])
            chunk = (
                (
                    await session.execute(
                        select(Chunk)
                        .where(Chunk.document_version_id == version_id)
                        .order_by(Chunk.chunk_index)
                        .limit(1)
                    )
                )
                .scalars()
                .first()
            )
            assert chunk is not None
            assert "OCR_FALLBACK_TEXT" in (chunk.text or "")

    asyncio.run(_verify_chunk_text_contains())


def test_docx_parser_ingestion_success(
    client: TestClient,
    unique_email: str,
) -> None:
    pytest.importorskip("docx")

    token = _register(client, f"docx_{unique_email}")
    ws = _create_workspace(client, token, "DOCXWS")

    from docx import Document as DocxDocument

    d = DocxDocument()
    d.add_paragraph("My Heading", style="Heading 1")
    d.add_paragraph("Body paragraph one", style=None)
    bio = BytesIO()
    d.save(bio)
    docx_bytes = bio.getvalue()

    doc_body, _storage_key, job_id = _create_document_via_upload_url(
        client,
        token=token,
        workspace_id=ws["id"],
        filename="sample.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_size=len(docx_bytes),
        content_bytes=docx_bytes,
    )

    process_job_sync(job_id)

    async def _verify() -> None:
        async with AsyncSessionLocal() as session:
            version_id = UUID(doc_body["latest_version_id"])
            chunks = (
                (
                    await session.execute(
                        select(Chunk)
                        .where(Chunk.document_version_id == version_id)
                        .order_by(Chunk.chunk_index)
                    )
                )
                .scalars()
                .all()
            )
            assert len(chunks) >= 1
            meta = chunks[0].metadata_json
            assert meta["section_heading"] == "My Heading"

    asyncio.run(_verify())


def test_pdf_native_ingestion_success(
    client: TestClient,
    unique_email: str,
) -> None:
    pytest.importorskip("fitz")

    import fitz

    token = _register(client, f"pdf_{unique_email}")
    ws = _create_workspace(client, token, "PDFWS")

    # Build a minimal PDF with one page and one text line.
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    page.insert_text((72, 72), "Hello PDF")
    bio = BytesIO()
    pdf_doc.save(bio)
    pdf_bytes = bio.getvalue()

    doc_body, _storage_key, job_id = _create_document_via_upload_url(
        client,
        token=token,
        workspace_id=ws["id"],
        filename="sample.pdf",
        content_type="application/pdf",
        file_size=len(pdf_bytes),
        content_bytes=pdf_bytes,
    )
    version_id = UUID(doc_body["latest_version_id"])

    process_job_sync(job_id)

    async def _verify() -> None:
        async with AsyncSessionLocal() as session:
            pages_count = (
                await session.execute(
                    select(func.count())
                    .select_from(Page)
                    .where(Page.document_version_id == version_id)
                )
            ).scalar_one()
            chunks = (
                (
                    await session.execute(
                        select(Chunk).where(Chunk.document_version_id == version_id)
                    )
                )
                .scalars()
                .all()
            )
            assert pages_count == 1
            assert len(chunks) >= 1

    asyncio.run(_verify())


def test_ingestion_failure_marks_job_and_document_failed(
    client: TestClient,
    unique_email: str,
) -> None:
    token = _register(client, f"fail_{unique_email}")
    ws = _create_workspace(client, token, "FAILWS")

    doc_body, storage_key, job_id = _create_document_via_upload_url(
        client,
        token=token,
        workspace_id=ws["id"],
        filename="bad.txt",
        content_type="text/plain",
        file_size=10,
        content_bytes=b"hello",
    )
    FAKE_OBJECT_RAISE_ON.add(storage_key)
    process_job_sync(job_id)
    FAKE_OBJECT_RAISE_ON.discard(storage_key)

    job_resp = client.get(f"/api/ingestion/jobs/{job_id}", headers=auth_headers(token))
    assert job_resp.status_code == 200
    assert job_resp.json()["job"]["status"] == "failed"
    assert job_resp.json()["job"]["error_code"] is not None


def test_reprocess_creates_new_job_and_resets_artifacts(
    client: TestClient,
    unique_email: str,
) -> None:
    token = _register(client, f"re_{unique_email}")
    ws = _create_workspace(client, token, "REPROCWS")

    doc_body, storage_key, old_job_id = _create_document_via_upload_url(
        client,
        token=token,
        workspace_id=ws["id"],
        filename="r.txt",
        content_type="text/plain",
        file_size=10,
        content_bytes=b"one\n\ntwo",
    )
    version_id = UUID(doc_body["latest_version_id"])

    process_job_sync(UUID(old_job_id))

    async def _counts() -> tuple[int, int]:
        async with AsyncSessionLocal() as session:
            pages_count = (
                await session.execute(
                    select(func.count())
                    .select_from(Page)
                    .where(Page.document_version_id == version_id)
                )
            ).scalar_one()
            chunks_count = (
                await session.execute(
                    select(func.count())
                    .select_from(Chunk)
                    .where(Chunk.document_version_id == version_id)
                )
            ).scalar_one()
            return int(pages_count), int(chunks_count)

    before_pages, before_chunks = asyncio.run(_counts())
    assert before_pages >= 1
    assert before_chunks >= 1

    resp = client.post(f"/api/documents/{doc_body['id']}/reprocess", headers=auth_headers(token))
    assert resp.status_code == 200
    re_body = resp.json()
    new_job_id = UUID(re_body["job"]["id"])
    assert new_job_id != UUID(old_job_id)
    assert re_body["job"]["status"] == "pending"
    assert re_body["job"]["stage"] == "queued"

    pages_after, chunks_after = asyncio.run(_counts())
    assert pages_after == 0
    assert chunks_after == 0

    process_job_sync(new_job_id)
    job_resp = client.get(f"/api/ingestion/jobs/{new_job_id}", headers=auth_headers(token))
    assert job_resp.status_code == 200
    assert job_resp.json()["job"]["status"] == "succeeded"


def test_ready_db_down(monkeypatch: pytest.MonkeyPatch) -> None:
    import redis.asyncio as redis_async

    class DummyConn:
        async def execute(self, _query: Any) -> None:
            return None

    class DummyConnectCM:
        async def __aenter__(self) -> DummyConn:
            raise Exception("db down")

        async def __aexit__(self, exc_type, exc, tb) -> None:
            return None

    class DummyEngine:
        def connect(self) -> DummyConnectCM:
            return DummyConnectCM()

    class DummyRedis:
        async def ping(self) -> bool:
            return True

        async def close(self) -> None:
            return None

    monkeypatch.setattr(health_mod, "engine", DummyEngine())
    monkeypatch.setattr(redis_async, "from_url", lambda *a, **k: DummyRedis())

    from app.main import app

    with TestClient(app) as c:
        resp = c.get("/ready")
        assert resp.status_code == 503
        body = resp.json()
        assert body["status"] == "not_ready"
        assert body["checks"]["db"] is False


def test_ready_redis_down(monkeypatch: pytest.MonkeyPatch) -> None:
    import redis.asyncio as redis_async

    class DummyConn:
        async def execute(self, _query: Any) -> None:
            return None

    class DummyConnectCM:
        async def __aenter__(self) -> DummyConn:
            return DummyConn()

        async def __aexit__(self, exc_type, exc, tb) -> None:
            return None

    class DummyEngine:
        def connect(self) -> DummyConnectCM:
            return DummyConnectCM()

    class DummyRedis:
        async def ping(self) -> bool:
            raise Exception("redis down")

        async def close(self) -> None:
            return None

    monkeypatch.setattr(health_mod, "engine", DummyEngine())
    monkeypatch.setattr(redis_async, "from_url", lambda *a, **k: DummyRedis())

    from app.main import app

    with TestClient(app) as c:
        resp = c.get("/ready")
        assert resp.status_code == 503
        body = resp.json()
        assert body["status"] == "not_ready"
        assert body["checks"]["redis"] is False
